import io
from typing import Any, Dict

# pyrefly: ignore [missing-import]
from docx import Document
# pyrefly: ignore [missing-import]
from docx.enum.table import WD_TABLE_ALIGNMENT
# pyrefly: ignore [missing-import]
from docx.enum.text import WD_ALIGN_PARAGRAPH
# pyrefly: ignore [missing-import]
from docx.oxml import OxmlElement, parse_xml
# pyrefly: ignore [missing-import]
from docx.oxml.ns import nsdecls, qn
# pyrefly: ignore [missing-import]
from docx.shared import Inches, Pt, RGBColor


def _set_cell_background(cell, color_hex: str):
    """Set the background color of a table cell in DOCX using XML shading."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tc_pr.append(shd)


def _set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner padding for table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def generate_docx_report(data: Dict[str, Any]) -> io.BytesIO:
    """Generate a professionally styled Word (DOCX) document for meeting intelligence."""
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Palette
    indigo = RGBColor(79, 70, 229)    # #4F46E5
    dark_slate = RGBColor(15, 23, 42) # #0F172A
    muted = RGBColor(100, 116, 139)   # #64748B
    green = RGBColor(5, 150, 105)     # #059669

    filename = data.get("filename", "Meeting Recording")
    date_str = data.get("date_str", "N/A")
    meeting_id = data.get("meeting_id", "N/A")

    # Document Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run(f"Meeting Report: {filename}")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = dark_slate

    # Metadata Paragraph
    meta_p = doc.add_paragraph()
    meta_p.paragraph_format.space_before = Pt(0)
    meta_p.paragraph_format.space_after = Pt(14)
    run_meta = meta_p.add_run(f"Date: {date_str}   |   Meeting ID: {meeting_id}")
    run_meta.font.name = "Calibri"
    run_meta.font.size = Pt(10)
    run_meta.font.color.rgb = muted

    # 1. Executive Summary
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(4)
    h1.paragraph_format.keep_with_next = True
    r_h1 = h1.add_run("EXECUTIVE SUMMARY")
    r_h1.font.name = "Calibri"
    r_h1.font.size = Pt(12)
    r_h1.font.bold = True
    r_h1.font.color.rgb = indigo

    summary_text = data.get("summary", "No executive summary available.")
    summary_p = doc.add_paragraph()
    summary_p.paragraph_format.space_before = Pt(4)
    summary_p.paragraph_format.space_after = Pt(12)
    summary_p.paragraph_format.line_spacing = 1.2
    r_sum = summary_p.add_run(summary_text)
    r_sum.font.name = "Calibri"
    r_sum.font.size = Pt(10.5)
    r_sum.font.color.rgb = dark_slate

    # 2. Key Discussion Points
    key_points = data.get("key_points", [])
    if key_points:
        h2 = doc.add_paragraph()
        h2.paragraph_format.space_before = Pt(16)
        h2.paragraph_format.space_after = Pt(4)
        h2.paragraph_format.keep_with_next = True
        r_h2 = h2.add_run("KEY DISCUSSION POINTS")
        r_h2.font.name = "Calibri"
        r_h2.font.size = Pt(12)
        r_h2.font.bold = True
        r_h2.font.color.rgb = indigo

        for pt in key_points:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(pt)
            r.font.name = "Calibri"
            r.font.size = Pt(10)
            r.font.color.rgb = dark_slate

    # 3. Decisions Made
    decisions = data.get("decisions", [])
    if decisions:
        h3 = doc.add_paragraph()
        h3.paragraph_format.space_before = Pt(16)
        h3.paragraph_format.space_after = Pt(4)
        h3.paragraph_format.keep_with_next = True
        r_h3 = h3.add_run("DECISIONS MADE")
        r_h3.font.name = "Calibri"
        r_h3.font.size = Pt(12)
        r_h3.font.bold = True
        r_h3.font.color.rgb = indigo

        for dec in decisions:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)

            r_tag = p.add_run("[AGREED] ")
            r_tag.font.name = "Calibri"
            r_tag.font.size = Pt(10)
            r_tag.font.bold = True
            r_tag.font.color.rgb = green

            r_dec = p.add_run(dec)
            r_dec.font.name = "Calibri"
            r_dec.font.size = Pt(10)
            r_dec.font.bold = True
            r_dec.font.color.rgb = dark_slate

    # 4. Action Items Table
    action_items = data.get("action_items", [])
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    h4.paragraph_format.keep_with_next = True
    r_h4 = h4.add_run(f"ACTION ITEMS ({len(action_items)})")
    r_h4.font.name = "Calibri"
    r_h4.font.size = Pt(12)
    r_h4.font.bold = True
    r_h4.font.color.rgb = indigo

    if not action_items:
        p = doc.add_paragraph()
        r = p.add_run("No action items were identified in this meeting.")
        r.font.name = "Calibri"
        r.font.size = Pt(10)
        r.font.italic = True
        r.font.color.rgb = muted
    else:
        table = doc.add_table(rows=len(action_items) + 1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        col_widths = [Inches(4.0), Inches(1.5), Inches(1.3)]
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = w

        # Header Row
        headers = ["TASK", "ASSIGNEE", "DEADLINE"]
        hdr_row = table.rows[0]
        # Repeat header on new pages
        tr_pr = hdr_row._tr.get_or_add_trPr()
        tr_pr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for i, title in enumerate(headers):
            cell = hdr_row.cells[i]
            _set_cell_background(cell, "4F46E5")
            _set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(title)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

        # Body Rows
        for idx, item in enumerate(action_items):
            row = table.rows[idx + 1]
            bg_color = "F8FAFC" if idx % 2 == 1 else "FFFFFF"

            task_cell = row.cells[0]
            assignee_cell = row.cells[1]
            deadline_cell = row.cells[2]

            for c in (task_cell, assignee_cell, deadline_cell):
                _set_cell_background(c, bg_color)
                _set_cell_margins(c, top=80, bottom=80, left=120, right=120)

            # Task
            p_task = task_cell.paragraphs[0]
            p_task.paragraph_format.space_before = Pt(0)
            p_task.paragraph_format.space_after = Pt(0)
            r_t = p_task.add_run(item.get("task", ""))
            r_t.font.name = "Calibri"
            r_t.font.size = Pt(9.5)
            r_t.font.color.rgb = dark_slate

            # Assignee
            p_ass = assignee_cell.paragraphs[0]
            p_ass.paragraph_format.space_before = Pt(0)
            p_ass.paragraph_format.space_after = Pt(0)
            assignee = item.get("assignee")
            r_a = p_ass.add_run(assignee if assignee else "Unassigned")
            r_a.font.name = "Calibri"
            r_a.font.size = Pt(9.5)
            if assignee:
                r_a.font.bold = True
                r_a.font.color.rgb = dark_slate
            else:
                r_a.font.italic = True
                r_a.font.color.rgb = muted

            # Deadline
            p_dl = deadline_cell.paragraphs[0]
            p_dl.paragraph_format.space_before = Pt(0)
            p_dl.paragraph_format.space_after = Pt(0)
            deadline = item.get("deadline")
            r_d = p_dl.add_run(deadline if deadline else "None")
            r_d.font.name = "Calibri"
            r_d.font.size = Pt(9.5)
            if deadline:
                r_d.font.color.rgb = dark_slate
            else:
                r_d.font.italic = True
                r_d.font.color.rgb = muted

    # Save into in-memory buffer
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
